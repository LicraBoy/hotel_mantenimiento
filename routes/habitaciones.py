from flask import Blueprint, render_template, request, redirect, session, flash, send_file
import io
import os
from datetime import date
from database.db import conectar
from werkzeug.security import check_password_hash
from routes.auth import requiere_rol
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment

habitaciones_bp = Blueprint("habitaciones", __name__)

ESTADOS_HAB_VALIDOS = {"Disponible", "En mantenimiento", "Ocupada", "Fuera de servicio"}


def _validar_habitacion(form):
    errores = []
    numero = form.get("numero", "").strip()
    estado = form.get("estado", "").strip()
    if not numero:
        errores.append("El número de habitación es obligatorio.")
    if estado not in ESTADOS_HAB_VALIDOS:
        errores.append("Estado de habitación no válido.")
    return errores


# ===============================
# Listar habitaciones
# ===============================
@habitaciones_bp.route("/habitaciones")
def lista():
    if "user" not in session:
        return redirect("/login")

    conn = conectar()
    cursor = conn.cursor()
    # Modificado para incluir el mayor riesgo predictivo de la habitación
    cursor.execute("""
        SELECT h.id, h.numero, h.estado, h.fecha_ultimo_mantenimiento,
               COALESCE((
                   SELECT nivel_riesgo 
                   FROM predicciones 
                   WHERE habitacion_id = h.id 
                   ORDER BY prob_30_dias DESC 
                   LIMIT 1
               ), 'BAJO') as nivel_riesgo
        FROM habitaciones h
        ORDER BY h.numero
    """)
    data = cursor.fetchall()
    cursor.close()
    conn.close()

    return render_template("habitaciones.html", habitaciones=data)


# ===============================
# Nueva habitación
# ===============================
@habitaciones_bp.route("/habitaciones/nueva", methods=["GET", "POST"])
def nueva():
    if "user" not in session:
        return redirect("/login")

    if request.method == "POST":
        errores = _validar_habitacion(request.form)
        if errores:
            for e in errores:
                flash(f"⚠️ {e}")
            return render_template("habitacion_form.html")

        numero = request.form["numero"]
        estado = request.form["estado"]

        conn = conectar()
        cursor = conn.cursor()

        cursor.execute("SELECT id FROM habitaciones WHERE numero=%s", (numero,))
        existe = cursor.fetchone()

        if existe:
            cursor.close()
            conn.close()
            flash("⚠️ La habitación ya está registrada.")
            return redirect("/habitaciones")

        cursor.execute("""
            INSERT INTO habitaciones(numero, estado, fecha_ultimo_mantenimiento)
            VALUES (%s, %s, %s)
        """, (numero, estado, ""))

        conn.commit()
        cursor.close()
        conn.close()
        return redirect("/habitaciones")

    return render_template("habitacion_form.html")


# ===============================
# Editar habitación
# ===============================
@habitaciones_bp.route("/habitaciones/editar/<int:id>", methods=["GET", "POST"])
def editar(id):
    if "user" not in session:
        return redirect("/login")

    conn = conectar()
    cursor = conn.cursor()

    cursor.execute("SELECT * FROM habitaciones WHERE id=%s", (id,))
    habitacion = cursor.fetchone()

    if request.method == "POST":
        numero = request.form["numero"]
        estado = request.form["estado"]

        cursor.execute("""
            UPDATE habitaciones SET numero=%s, estado=%s WHERE id=%s
        """, (numero, estado, id))

        cursor.execute("""
            INSERT INTO historial_habitaciones
            (habitacion_id, accion, descripcion, usuario, fecha)
            VALUES (%s, %s, %s, %s, NOW())
        """, (id, "EDICION", "Actualización de habitación", session["user"]))

        conn.commit()
        cursor.close()
        conn.close()
        return redirect("/habitaciones")

    cursor.close()
    conn.close()
    return render_template("habitacion_editar.html", habitacion=habitacion)


# ===============================
# Eliminar habitación
# ===============================
@habitaciones_bp.route("/habitaciones/eliminar/<int:id>", methods=["POST"])
@requiere_rol("admin")
def eliminar(id):
    admin_pw = request.form.get("admin_password")
    user_id = session.get("user_id")

    if not admin_pw:
        flash("⚠️ Debes proporcionar la contraseña de administrador.")
        return redirect("/habitaciones")

    conn = conectar()
    cursor = conn.cursor()

    cursor.execute("SELECT password FROM usuarios WHERE id=%s", (user_id,))
    row = cursor.fetchone()

    if not row or not check_password_hash(row[0], admin_pw):
        cursor.close()
        conn.close()
        flash("❌ Contraseña incorrecta. Eliminación cancelada.")
        return redirect("/habitaciones")

    cursor.execute("DELETE FROM habitaciones WHERE id=%s", (id,))
    conn.commit()
    cursor.close()
    conn.close()
    
    flash("✅ Habitación eliminada con éxito.")
    return redirect("/habitaciones")


# ===============================
# Detalle de habitación
# ===============================
@habitaciones_bp.route("/habitacion/<numero>")
def detalle(numero):
    if "user" not in session:
        return redirect("/login")

    conn = conectar()
    cursor = conn.cursor()

    cursor.execute("""
        SELECT id, numero, estado FROM habitaciones WHERE numero=%s
    """, (numero,))
    habitacion_row = cursor.fetchone()

    if not habitacion_row:
        cursor.close()
        conn.close()
        return "Habitación no encontrada"

    habitacion_id = habitacion_row[0]
    habitacion = (habitacion_row[1], habitacion_row[2])

    # Línea de tiempo cronológica
    cursor.execute("""
        SELECT fecha::text, 'Inspección' as tipo, observaciones as detalle,
               CASE WHEN genera_mantenimiento = 1
                    THEN 'En mantenimiento' ELSE 'Disponible'
               END as estado_final
        FROM inspecciones WHERE habitacion_id=%s

        UNION ALL

        SELECT fecha::text, 'Mantenimiento' as tipo,
               COALESCE(elemento, '') || ' — ' || COALESCE(descripcion, '') as detalle,
               COALESCE(estado, 'Finalizado') as estado_final
        FROM mantenimiento WHERE habitacion_id=%s

        ORDER BY fecha DESC
    """, (habitacion_id, habitacion_id))
    timeline = cursor.fetchall()

    # Total gastado
    cursor.execute("""
        SELECT COALESCE(SUM(costo),0) FROM mantenimiento WHERE habitacion_id=%s
    """, (habitacion_id,))
    total_gastado = cursor.fetchone()[0]

    # Número total de fallas
    cursor.execute("SELECT COUNT(*) FROM mantenimiento WHERE habitacion_id=%s", (habitacion_id,))
    total_fallas = cursor.fetchone()[0]

    # Última fecha de mantenimiento
    cursor.execute("""
        SELECT fecha FROM mantenimiento
        WHERE habitacion_id=%s ORDER BY fecha DESC LIMIT 1
    """, (habitacion_id,))
    ultima_fecha = cursor.fetchone()
    ultima_fecha = ultima_fecha[0] if ultima_fecha else "Sin registros"

    cursor.close()
    conn.close()

    return render_template(
        "detalle_habitacion.html",
        habitacion=habitacion,
        habitacion_id=habitacion_id,
        timeline=timeline,
        total_gastado=total_gastado,
        total_fallas=total_fallas,
        ultima_fecha=ultima_fecha
    )


# ===============================
# Exportar lista de habitaciones
# ===============================
@habitaciones_bp.route("/habitaciones/exportar")
def exportar():
    if "user" not in session:
        return redirect("/login")

    conn = conectar()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT h.numero, h.estado, h.fecha_ultimo_mantenimiento,
               COUNT(m.id) as total_fallas,
               COALESCE(SUM(m.costo), 0) as total_gastado,
               COALESCE(MAX(p.nivel_riesgo), 'BAJO') as nivel_riesgo
        FROM habitaciones h
        LEFT JOIN mantenimiento m ON m.habitacion_id = h.id
        LEFT JOIN predicciones p ON p.habitacion_id = h.id
        GROUP BY h.id, h.numero, h.estado, h.fecha_ultimo_mantenimiento
        ORDER BY h.numero
    """)
    rows = cursor.fetchall()
    cursor.close()
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Habitaciones"

    encabezados = ["Número", "Estado", "Último Mantenimiento", "Total Fallas", "Costo Total ($)", "Riesgo IA"]
    header_fill = PatternFill("solid", fgColor="1E3A5F")
    header_font = Font(color="FFFFFF", bold=True)

    for col, titulo in enumerate(encabezados, start=1):
        cell = ws.cell(row=1, column=col, value=titulo)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row in rows:
        ws.append([
            str(row[0]),
            row[1],
            str(row[2]) if row[2] else "Sin registro",
            int(row[3]),
            float(row[4]),
            row[5],
        ])

    anchos = [12, 20, 22, 14, 18, 12]
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = ancho

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nombre = f"habitaciones_{date.today()}.xlsx"
    return send_file(buffer, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True, download_name=nombre)


# ===============================
# Exportar historial de habitación
# ===============================
@habitaciones_bp.route("/habitacion/<numero>/exportar")
def exportar_historial(numero):
    if "user" not in session:
        return redirect("/login")

    conn = conectar()
    cursor = conn.cursor()

    cursor.execute("SELECT id FROM habitaciones WHERE numero=%s", (numero,))
    row = cursor.fetchone()
    if not row:
        cursor.close()
        conn.close()
        return "Habitación no encontrada", 404

    habitacion_id = row[0]

    cursor.execute("""
        SELECT fecha::text, 'Inspección' as tipo, observaciones as detalle,
               CASE WHEN genera_mantenimiento = 1 THEN 'En mantenimiento' ELSE 'Disponible' END as estado_final
        FROM inspecciones WHERE habitacion_id=%s
        UNION ALL
        SELECT fecha::text, 'Mantenimiento' as tipo,
               COALESCE(elemento, '') || ' — ' || COALESCE(descripcion, '') as detalle,
               COALESCE(estado, 'Finalizado') as estado_final
        FROM mantenimiento WHERE habitacion_id=%s
        ORDER BY fecha DESC
    """, (habitacion_id, habitacion_id))
    timeline = cursor.fetchall()
    cursor.close()
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = f"Habitación {numero}"

    encabezados = ["Fecha", "Tipo", "Detalle", "Estado Final"]
    header_fill = PatternFill("solid", fgColor="1E3A5F")
    header_font = Font(color="FFFFFF", bold=True)

    for col, titulo in enumerate(encabezados, start=1):
        cell = ws.cell(row=1, column=col, value=titulo)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row in timeline:
        ws.append([row[0], row[1], row[2] or "", row[3]])

    for i, ancho in enumerate([16, 14, 50, 20], start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = ancho

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nombre = f"habitacion_{numero}_historial_{date.today()}.xlsx"
    return send_file(buffer, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True, download_name=nombre)
