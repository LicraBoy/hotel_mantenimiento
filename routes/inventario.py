"""
Blueprint: Inventario General de Activos
Catálogo maestro de todos los activos del hotel.
"""
import os
import io
from datetime import date, datetime
from flask import (Blueprint, render_template, request, redirect,
                   session, flash, send_file, jsonify, url_for)
from database.db import conectar
from routes.auth import requiere_login, requiere_rol
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment

inventario_bp = Blueprint("inventario", __name__,
                          template_folder="../templates/inventario")

ESTADOS_ACTIVO = ["operativo", "en_mantenimiento", "fuera_servicio", "dado_de_baja"]
TIPOS_DOCUMENTO = ["manual", "garantia", "factura", "contrato", "otro"]
UPLOAD_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "static", "uploads", "activos")
QR_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "static", "uploads", "qr")
DOC_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "static", "uploads", "documentos")

# Siglas para código automático
SIGLAS_CATEGORIA = {
    "Línea Blanca": "LB", "Climatización": "CL", "Electricidad": "EL",
    "Plomería": "PL", "Mobiliario": "MO", "Tecnología": "TE",
    "Seguridad": "SE", "Cocina": "CO", "Elevación": "EV", "Otros": "OT",
}


def _generar_codigo(cursor, categoria_nombre):
    """Auto-genera código: [SIGLAS]-[AÑO]-[SEC 3 dígitos]"""
    siglas = SIGLAS_CATEGORIA.get(categoria_nombre, "XX")
    anio = date.today().year
    prefijo = f"{siglas}-{anio}-"
    cursor.execute(
        "SELECT codigo FROM activos WHERE codigo LIKE %s ORDER BY codigo DESC LIMIT 1",
        (f"{prefijo}%",)
    )
    ultimo = cursor.fetchone()
    if ultimo:
        try:
            sec = int(ultimo[0].split("-")[-1]) + 1
        except (ValueError, IndexError):
            sec = 1
    else:
        sec = 1
    return f"{prefijo}{sec:03d}"


def _guardar_foto(foto, activo_id):
    """Guarda foto del activo como webp."""
    os.makedirs(UPLOAD_DIR, exist_ok=True)
    from PIL import Image
    nombre = f"activo_{activo_id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.webp"
    ruta = os.path.join(UPLOAD_DIR, nombre)
    try:
        img = Image.open(foto)
        if img.mode in ("RGBA", "P"):
            img = img.convert("RGB")
        img.thumbnail((1080, 1080), Image.Resampling.LANCZOS)
        img.save(ruta, "WEBP", quality=80)
        return f"/static/uploads/activos/{nombre}"
    except Exception:
        return None


def _generar_qr(activo_id, codigo):
    """Genera QR con la URL de escaneo del activo."""
    os.makedirs(QR_DIR, exist_ok=True)
    import qrcode
    url = f"/inventario/escanear/{codigo}"
    img = qrcode.make(url)
    nombre = f"activo_{activo_id}.png"
    ruta = os.path.join(QR_DIR, nombre)
    img.save(ruta)
    return f"/static/uploads/qr/{nombre}"


def _estado_garantia(garantia_hasta):
    """Retorna ('vigente'|'por_vencer'|'vencida', dias_restantes)"""
    if not garantia_hasta:
        return None, None
    hoy = date.today()
    if isinstance(garantia_hasta, str):
        try:
            garantia_hasta = datetime.strptime(garantia_hasta, "%Y-%m-%d").date()
        except ValueError:
            return None, None
    dias = (garantia_hasta - hoy).days
    if dias < 0:
        return "vencida", dias
    elif dias <= 90:
        return "por_vencer", dias
    return "vigente", dias


# ═══════════════════════════════════════════
# RUTAS
# ═══════════════════════════════════════════

@inventario_bp.route("/inventario")
@requiere_login
def index():
    conn = conectar()
    cursor = conn.cursor()

    # Categorías con conteo
    cursor.execute("""
        SELECT c.id, c.nombre, c.icono, c.color, c.descripcion,
               COUNT(a.id) as total
        FROM categorias_activos c
        LEFT JOIN activos a ON a.categoria_id = c.id
        WHERE c.activo = TRUE
        GROUP BY c.id, c.nombre, c.icono, c.color, c.descripcion
        ORDER BY c.nombre
    """)
    categorias = cursor.fetchall()

    # Estadísticas
    cursor.execute("SELECT COUNT(*) FROM activos")
    total = cursor.fetchone()[0]

    cursor.execute("SELECT COUNT(*) FROM activos WHERE estado='operativo'")
    operativos = cursor.fetchone()[0]

    cursor.execute("SELECT COUNT(*) FROM activos WHERE estado='en_mantenimiento'")
    en_mant = cursor.fetchone()[0]

    cursor.execute("SELECT COUNT(*) FROM activos WHERE estado='fuera_servicio'")
    fuera = cursor.fetchone()[0]

    cursor.execute("SELECT COALESCE(SUM(costo_adquisicion), 0) FROM activos")
    valor_total = cursor.fetchone()[0]

    cursor.execute("SELECT COUNT(*) FROM activos WHERE garantia_hasta < CURRENT_DATE")
    garantias_vencidas = cursor.fetchone()[0]

    cursor.close()
    conn.close()

    return render_template("inventario/index.html",
                           categorias=categorias,
                           total=total, operativos=operativos,
                           en_mant=en_mant, fuera=fuera,
                           valor_total=valor_total,
                           garantias_vencidas=garantias_vencidas)


@inventario_bp.route("/inventario/lista")
@requiere_login
def lista():
    filtro_cat = request.args.get("categoria", "")
    filtro_estado = request.args.get("estado", "")
    filtro_busq = request.args.get("q", "").strip()
    pagina = int(request.args.get("pagina", 1))
    por_pagina = 25

    condiciones = []
    params = []

    if filtro_cat:
        condiciones.append("a.categoria_id = %s")
        params.append(int(filtro_cat))
    if filtro_estado:
        condiciones.append("a.estado = %s")
        params.append(filtro_estado)
    if filtro_busq:
        condiciones.append("(a.nombre ILIKE %s OR a.codigo ILIKE %s OR a.marca ILIKE %s OR a.numero_serie ILIKE %s)")
        like = f"%{filtro_busq}%"
        params.extend([like, like, like, like])

    where = ("WHERE " + " AND ".join(condiciones)) if condiciones else ""

    conn = conectar()
    cursor = conn.cursor()

    cursor.execute(f"SELECT COUNT(*) FROM activos a {where}", params)
    total = cursor.fetchone()[0]
    total_pags = max(1, (total + por_pagina - 1) // por_pagina)
    offset = (pagina - 1) * por_pagina

    cursor.execute(f"""
        SELECT a.id, a.codigo, a.nombre, c.nombre, c.color,
               a.marca, a.modelo, a.ubicacion, a.estado,
               a.garantia_hasta, a.costo_adquisicion, a.foto_url
        FROM activos a
        LEFT JOIN categorias_activos c ON a.categoria_id = c.id
        {where}
        ORDER BY a.codigo
        LIMIT %s OFFSET %s
    """, params + [por_pagina, offset])
    activos = cursor.fetchall()

    cursor.execute("SELECT id, nombre FROM categorias_activos WHERE activo=TRUE ORDER BY nombre")
    categorias = cursor.fetchall()

    cursor.close()
    conn.close()

    # Enriquecer con estado de garantía
    activos_enriq = []
    for a in activos:
        estado_gar, dias = _estado_garantia(a[9])
        activos_enriq.append(list(a) + [estado_gar, dias])

    return render_template("inventario/lista.html",
                           activos=activos_enriq, categorias=categorias,
                           filtro_cat=filtro_cat, filtro_estado=filtro_estado,
                           filtro_busq=filtro_busq, estados=ESTADOS_ACTIVO,
                           pagina=pagina, total_pags=total_pags, total=total)


@inventario_bp.route("/inventario/activo/<int:id>")
@requiere_login
def detalle(id):
    conn = conectar()
    cursor = conn.cursor()

    cursor.execute("""
        SELECT a.*, c.nombre as cat_nombre, c.icono, c.color,
               h.numero as hab_numero, u.username as creado_nombre
        FROM activos a
        LEFT JOIN categorias_activos c ON a.categoria_id = c.id
        LEFT JOIN habitaciones h ON a.habitacion_id = h.id
        LEFT JOIN usuarios u ON a.creado_por = u.id
        WHERE a.id = %s
    """, (id,))
    activo = cursor.fetchone()

    if not activo:
        cursor.close()
        conn.close()
        flash("⚠️ Activo no encontrado.")
        return redirect("/inventario")

    # Columnas: 0=id, 1=codigo, 2=nombre, 3=categoria_id, 4=subcategoria,
    # 5=marca, 6=modelo, 7=numero_serie, 8=fecha_compra, 9=fecha_instalacion,
    # 10=garantia_hasta, 11=proveedor, 12=costo_adquisicion, 13=vida_util_anos,
    # 14=ubicacion, 15=habitacion_id, 16=estado, 17=foto_url, 18=codigo_qr,
    # 19=notas, 20=creado_por, 21=creado_en, 22=actualizado_en
    # 23=cat_nombre, 24=icono, 25=color, 26=hab_numero, 27=creado_nombre

    estado_gar, dias_gar = _estado_garantia(activo[10])

    # Mantenimientos vinculados
    cursor.execute("""
        SELECT m.id, m.fecha, m.tipo, m.elemento, m.descripcion,
               m.tecnico, m.estado, m.costo, m.prioridad
        FROM mantenimiento m
        WHERE m.activo_id = %s
        ORDER BY m.fecha DESC LIMIT 50
    """, (id,))
    mantenimientos = cursor.fetchall()

    # Documentos
    cursor.execute("""
        SELECT id, tipo, nombre_archivo, ruta_archivo, subido_en
        FROM documentos_activo
        WHERE activo_id = %s ORDER BY subido_en DESC
    """, (id,))
    documentos = cursor.fetchall()

    cursor.close()
    conn.close()

    return render_template("inventario/detalle.html",
                           a=activo, estado_gar=estado_gar,
                           dias_gar=dias_gar,
                           mantenimientos=mantenimientos,
                           documentos=documentos)


@inventario_bp.route("/inventario/activo/nuevo", methods=["GET", "POST"])
@requiere_login
def nuevo():
    conn = conectar()
    cursor = conn.cursor()

    if request.method == "POST":
        nombre = request.form.get("nombre", "").strip()
        categoria_id = request.form.get("categoria_id")
        codigo = request.form.get("codigo", "").strip()
        subcategoria = request.form.get("subcategoria", "").strip()
        marca = request.form.get("marca", "").strip()
        modelo = request.form.get("modelo", "").strip()
        numero_serie = request.form.get("numero_serie", "").strip()
        fecha_compra = request.form.get("fecha_compra") or None
        fecha_instalacion = request.form.get("fecha_instalacion") or None
        garantia_hasta = request.form.get("garantia_hasta") or None
        proveedor = request.form.get("proveedor", "").strip()
        costo_adquisicion = request.form.get("costo_adquisicion") or 0
        vida_util_anos = request.form.get("vida_util_anos") or None
        ubicacion = request.form.get("ubicacion", "").strip()
        habitacion_id = request.form.get("habitacion_id") or None
        estado = request.form.get("estado", "operativo")
        notas = request.form.get("notas", "").strip()

        if not nombre:
            flash("⚠️ El nombre del activo es obligatorio.")
            cursor.execute("SELECT id, nombre FROM categorias_activos WHERE activo=TRUE ORDER BY nombre")
            cats = cursor.fetchall()
            cursor.execute("SELECT id, numero FROM habitaciones ORDER BY numero")
            habs = cursor.fetchall()
            cursor.close()
            conn.close()
            return render_template("inventario/form_activo.html",
                                   categorias=cats, habitaciones=habs,
                                   estados=ESTADOS_ACTIVO, modo="nuevo")

        # Auto-generar código si vacío
        if not codigo and categoria_id:
            cursor.execute("SELECT nombre FROM categorias_activos WHERE id=%s", (int(categoria_id),))
            cat_row = cursor.fetchone()
            cat_nombre = cat_row[0] if cat_row else "Otros"
            codigo = _generar_codigo(cursor, cat_nombre)

        cursor.execute("""
            INSERT INTO activos
            (codigo, nombre, categoria_id, subcategoria, marca, modelo,
             numero_serie, fecha_compra, fecha_instalacion, garantia_hasta,
             proveedor, costo_adquisicion, vida_util_anos, ubicacion,
             habitacion_id, estado, notas, creado_por)
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
            RETURNING id
        """, (codigo, nombre, int(categoria_id) if categoria_id else None,
              subcategoria, marca, modelo, numero_serie,
              fecha_compra, fecha_instalacion, garantia_hasta,
              proveedor, float(costo_adquisicion), int(vida_util_anos) if vida_util_anos else None,
              ubicacion, int(habitacion_id) if habitacion_id else None,
              estado, notas, session.get("user_id")))
        activo_id = cursor.fetchone()[0]

        # Foto
        foto = request.files.get("foto")
        if foto and foto.filename:
            foto_url = _guardar_foto(foto, activo_id)
            if foto_url:
                cursor.execute("UPDATE activos SET foto_url=%s WHERE id=%s", (foto_url, activo_id))

        # QR
        qr_url = _generar_qr(activo_id, codigo)
        cursor.execute("UPDATE activos SET codigo_qr=%s WHERE id=%s", (qr_url, activo_id))

        conn.commit()
        cursor.close()
        conn.close()
        flash(f"✅ Activo '{nombre}' registrado con código {codigo}")
        return redirect(f"/inventario/activo/{activo_id}")

    # GET
    cursor.execute("SELECT id, nombre FROM categorias_activos WHERE activo=TRUE ORDER BY nombre")
    cats = cursor.fetchall()
    cursor.execute("SELECT id, numero FROM habitaciones ORDER BY numero")
    habs = cursor.fetchall()
    cursor.close()
    conn.close()

    return render_template("inventario/form_activo.html",
                           categorias=cats, habitaciones=habs,
                           estados=ESTADOS_ACTIVO, modo="nuevo")


@inventario_bp.route("/inventario/activo/<int:id>/editar", methods=["GET", "POST"])
@requiere_login
def editar(id):
    conn = conectar()
    cursor = conn.cursor()

    if request.method == "POST":
        nombre = request.form.get("nombre", "").strip()
        categoria_id = request.form.get("categoria_id") or None
        subcategoria = request.form.get("subcategoria", "").strip()
        marca = request.form.get("marca", "").strip()
        modelo = request.form.get("modelo", "").strip()
        numero_serie = request.form.get("numero_serie", "").strip()
        fecha_compra = request.form.get("fecha_compra") or None
        fecha_instalacion = request.form.get("fecha_instalacion") or None
        garantia_hasta = request.form.get("garantia_hasta") or None
        proveedor = request.form.get("proveedor", "").strip()
        costo_adquisicion = request.form.get("costo_adquisicion") or 0
        vida_util_anos = request.form.get("vida_util_anos") or None
        ubicacion = request.form.get("ubicacion", "").strip()
        habitacion_id = request.form.get("habitacion_id") or None
        estado = request.form.get("estado", "operativo")
        notas = request.form.get("notas", "").strip()

        cursor.execute("""
            UPDATE activos SET
                nombre=%s, categoria_id=%s, subcategoria=%s, marca=%s, modelo=%s,
                numero_serie=%s, fecha_compra=%s, fecha_instalacion=%s,
                garantia_hasta=%s, proveedor=%s, costo_adquisicion=%s,
                vida_util_anos=%s, ubicacion=%s, habitacion_id=%s,
                estado=%s, notas=%s, actualizado_en=NOW()
            WHERE id=%s
        """, (nombre, int(categoria_id) if categoria_id else None,
              subcategoria, marca, modelo, numero_serie,
              fecha_compra, fecha_instalacion, garantia_hasta,
              proveedor, float(costo_adquisicion),
              int(vida_util_anos) if vida_util_anos else None,
              ubicacion, int(habitacion_id) if habitacion_id else None,
              estado, notas, id))

        foto = request.files.get("foto")
        if foto and foto.filename:
            foto_url = _guardar_foto(foto, id)
            if foto_url:
                cursor.execute("UPDATE activos SET foto_url=%s WHERE id=%s", (foto_url, id))

        conn.commit()
        cursor.close()
        conn.close()
        flash("✅ Activo actualizado correctamente.")
        return redirect(f"/inventario/activo/{id}")

    # GET
    cursor.execute("SELECT * FROM activos WHERE id=%s", (id,))
    activo = cursor.fetchone()
    if not activo:
        cursor.close()
        conn.close()
        flash("⚠️ Activo no encontrado.")
        return redirect("/inventario")

    cursor.execute("SELECT id, nombre FROM categorias_activos WHERE activo=TRUE ORDER BY nombre")
    cats = cursor.fetchall()
    cursor.execute("SELECT id, numero FROM habitaciones ORDER BY numero")
    habs = cursor.fetchall()
    cursor.close()
    conn.close()

    return render_template("inventario/form_activo.html",
                           categorias=cats, habitaciones=habs,
                           estados=ESTADOS_ACTIVO, modo="editar", activo=activo)


@inventario_bp.route("/inventario/activo/<int:id>/estado", methods=["POST"])
@requiere_login
def cambiar_estado(id):
    nuevo_estado = request.json.get("estado", "")
    if nuevo_estado not in ESTADOS_ACTIVO:
        return jsonify({"error": "Estado no válido"}), 400

    conn = conectar()
    cursor = conn.cursor()
    cursor.execute("UPDATE activos SET estado=%s, actualizado_en=NOW() WHERE id=%s",
                   (nuevo_estado, id))
    conn.commit()
    cursor.close()
    conn.close()
    return jsonify({"ok": True, "estado": nuevo_estado})


@inventario_bp.route("/inventario/activo/<int:id>/qr")
@requiere_login
def descargar_qr(id):
    conn = conectar()
    cursor = conn.cursor()
    cursor.execute("SELECT codigo, codigo_qr FROM activos WHERE id=%s", (id,))
    row = cursor.fetchone()
    if not row or not row[1]:
        # Regenerar
        cursor.execute("SELECT codigo FROM activos WHERE id=%s", (id,))
        r2 = cursor.fetchone()
        if r2:
            qr_url = _generar_qr(id, r2[0])
            cursor.execute("UPDATE activos SET codigo_qr=%s WHERE id=%s", (qr_url, id))
            conn.commit()
            row = (r2[0], qr_url)
    cursor.close()
    conn.close()

    if row and row[1]:
        ruta = os.path.join(os.path.dirname(os.path.dirname(__file__)), row[1].lstrip("/"))
        if os.path.exists(ruta):
            return send_file(ruta, mimetype="image/png", as_attachment=True,
                             download_name=f"QR_{row[0]}.png")
    flash("⚠️ No se pudo generar el QR.")
    return redirect(f"/inventario/activo/{id}")


@inventario_bp.route("/inventario/activo/<int:id>/documento", methods=["POST"])
@requiere_login
def subir_documento(id):
    os.makedirs(DOC_DIR, exist_ok=True)
    tipo = request.form.get("tipo_documento", "otro")
    archivo = request.files.get("documento")
    if not archivo or not archivo.filename:
        flash("⚠️ Selecciona un archivo.")
        return redirect(f"/inventario/activo/{id}")

    ext = archivo.filename.rsplit(".", 1)[-1].lower() if "." in archivo.filename else ""
    permitidas = {"pdf", "jpg", "jpeg", "png", "webp", "doc", "docx", "xls", "xlsx"}
    if ext not in permitidas:
        flash(f"⚠️ Extensión .{ext} no permitida.")
        return redirect(f"/inventario/activo/{id}")

    nombre = f"doc_{id}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.{ext}"
    ruta = os.path.join(DOC_DIR, nombre)
    archivo.save(ruta)
    ruta_web = f"/static/uploads/documentos/{nombre}"

    conn = conectar()
    cursor = conn.cursor()
    cursor.execute("""
        INSERT INTO documentos_activo (activo_id, tipo, nombre_archivo, ruta_archivo)
        VALUES (%s, %s, %s, %s)
    """, (id, tipo, archivo.filename, ruta_web))
    conn.commit()
    cursor.close()
    conn.close()

    flash(f"✅ Documento '{archivo.filename}' subido correctamente.")
    return redirect(f"/inventario/activo/{id}")


@inventario_bp.route("/inventario/categorias")
@requiere_rol("admin")
def categorias():
    conn = conectar()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM categorias_activos ORDER BY nombre")
    cats = cursor.fetchall()
    cursor.execute("""
        SELECT c.id, COUNT(a.id)
        FROM categorias_activos c
        LEFT JOIN activos a ON a.categoria_id = c.id
        GROUP BY c.id
    """)
    conteos = dict(cursor.fetchall())
    cursor.close()
    conn.close()
    return render_template("inventario/categorias.html", categorias=cats, conteos=conteos)


@inventario_bp.route("/inventario/categorias/nueva", methods=["POST"])
@requiere_rol("admin")
def categoria_nueva():
    nombre = request.form.get("nombre", "").strip()
    icono = request.form.get("icono", "fas fa-box").strip()
    color = request.form.get("color", "#6B7280").strip()
    descripcion = request.form.get("descripcion", "").strip()
    if not nombre:
        flash("⚠️ El nombre es obligatorio.")
        return redirect("/inventario/categorias")
    conn = conectar()
    cursor = conn.cursor()
    cursor.execute("""
        INSERT INTO categorias_activos (nombre, icono, color, descripcion)
        VALUES (%s, %s, %s, %s)
    """, (nombre, icono, color, descripcion))
    conn.commit()
    cursor.close()
    conn.close()
    flash(f"✅ Categoría '{nombre}' creada.")
    return redirect("/inventario/categorias")


@inventario_bp.route("/inventario/exportar")
@requiere_login
def exportar():
    conn = conectar()
    cursor = conn.cursor()

    cursor.execute("""
        SELECT a.codigo, a.nombre, c.nombre, a.marca, a.modelo,
               a.numero_serie, a.ubicacion, a.estado, a.fecha_compra,
               a.garantia_hasta, a.costo_adquisicion, a.vida_util_anos
        FROM activos a
        LEFT JOIN categorias_activos c ON a.categoria_id = c.id
        ORDER BY c.nombre, a.codigo
    """)
    rows = cursor.fetchall()
    cursor.close()
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario General"

    enc = ["Código", "Nombre", "Categoría", "Marca", "Modelo", "N° Serie",
           "Ubicación", "Estado", "Fecha Compra", "Garantía Hasta", "Costo ($)", "Vida Útil (años)"]
    hfill = PatternFill("solid", fgColor="1E3A5F")
    hfont = Font(color="FFFFFF", bold=True)

    for col, t in enumerate(enc, 1):
        cell = ws.cell(row=1, column=col, value=t)
        cell.fill = hfill
        cell.font = hfont
        cell.alignment = Alignment(horizontal="center")

    for row in rows:
        ws.append([
            row[0], row[1], row[2] or "", row[3] or "", row[4] or "",
            row[5] or "", row[6] or "", row[7] or "",
            str(row[8]) if row[8] else "", str(row[9]) if row[9] else "",
            float(row[10]) if row[10] else 0.0,
            row[11] if row[11] else "",
        ])

    anchos = [14, 25, 16, 14, 14, 16, 20, 16, 14, 14, 12, 12]
    for i, a in enumerate(anchos, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = a

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True, download_name=f"inventario_{date.today()}.xlsx")


@inventario_bp.route("/inventario/escanear/<codigo>")
@requiere_login
def escanear(codigo):
    conn = conectar()
    cursor = conn.cursor()
    cursor.execute("SELECT id FROM activos WHERE codigo=%s", (codigo,))
    row = cursor.fetchone()
    cursor.close()
    conn.close()
    if row:
        return redirect(f"/inventario/activo/{row[0]}")
    flash(f"⚠️ No se encontró activo con código '{codigo}'.")
    return redirect("/inventario")
