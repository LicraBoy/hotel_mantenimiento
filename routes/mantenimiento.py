from flask import Blueprint, render_template, request, redirect, session, flash, send_file
import io
from datetime import datetime, date
from database.db import conectar
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment

mantenimiento_bp = Blueprint("mantenimiento", __name__)

TIPOS_MANT_VALIDOS = {"Correctivo", "Preventivo", "Emergencia"}
ESTADOS_MANT_VALIDOS = {"Pendiente", "En proceso", "Completado"}


def _validar_mantenimiento(form):
    errores = []
    tipo = form.get("tipo", "").strip()
    elemento = form.get("elemento", "").strip()
    fecha = form.get("fecha", "").strip()
    costo = form.get("costo", "").strip()

    if tipo not in TIPOS_MANT_VALIDOS:
        errores.append("Tipo de mantenimiento no válido.")
    if not elemento or len(elemento) > 200:
        errores.append("El elemento es obligatorio (máximo 200 caracteres).")
    if not fecha:
        errores.append("La fecha es obligatoria.")
    else:
        try:
            datetime.strptime(fecha, "%Y-%m-%d")
        except ValueError:
            errores.append("Formato de fecha inválido (use AAAA-MM-DD).")
    if costo:
        try:
            if float(costo) < 0:
                errores.append("El costo no puede ser negativo.")
        except ValueError:
            errores.append("El costo debe ser un número.")
    return errores


# ===============================
# Listar mantenimientos
# ===============================
@mantenimiento_bp.route("/mantenimientos")
def lista():
    if "user" not in session:
        return redirect("/login")

    filtro_estado     = request.args.get("estado", "")
    filtro_tipo       = request.args.get("tipo", "")
    filtro_habitacion = request.args.get("habitacion", "")

    condiciones = []
    params = []

    if filtro_estado:
        condiciones.append("m.estado = %s")
        params.append(filtro_estado)
    else:
        condiciones.append("(m.estado != 'Completado' OR m.estado IS NULL)")

    if filtro_tipo:
        condiciones.append("m.tipo = %s")
        params.append(filtro_tipo)

    if filtro_habitacion:
        condiciones.append("h.numero ILIKE %s")
        params.append(f"%{filtro_habitacion}%")

    where = "WHERE " + " AND ".join(condiciones)

    conn = conectar()
    cursor = conn.cursor()
    cursor.execute(f"""
        SELECT h.numero, m.id, m.fecha, m.tipo, m.elemento,
               m.descripcion, m.tecnico, m.estado, m.costo
        FROM mantenimiento m
        JOIN habitaciones h ON m.habitacion_id = h.id
        {where}
        ORDER BY h.numero, m.fecha DESC
    """, params)
    rows = cursor.fetchall()
    cursor.close()
    conn.close()

    habitaciones_mant = {}
    for row in rows:
        numero = row[0]
        if numero not in habitaciones_mant:
            habitaciones_mant[numero] = []
        habitaciones_mant[numero].append(row[1:])

    habitaciones_ordenadas = sorted(habitaciones_mant.items(), key=lambda x: str(x[0]))

    return render_template(
        "mantenimientos.html",
        habitaciones_mant=habitaciones_ordenadas,
        filtro_estado=filtro_estado,
        filtro_tipo=filtro_tipo,
        filtro_habitacion=filtro_habitacion,
        tipos=sorted(TIPOS_MANT_VALIDOS),
        estados=sorted(ESTADOS_MANT_VALIDOS),
    )


# ===============================
# Nuevo mantenimiento
# ===============================
@mantenimiento_bp.route("/mantenimiento/nuevo/<int:habitacion_id>", methods=["GET", "POST"])
def nuevo(habitacion_id):
    if "user" not in session:
        return redirect("/login")

    if request.method == "POST":
        errores = _validar_mantenimiento(request.form)
        if errores:
            for e in errores:
                flash(f"⚠️ {e}")
            deteccion_id_get = request.form.get("deteccion_id", "")
            return render_template("mantenimiento_form.html", habitacion_id=habitacion_id, deteccion_id=deteccion_id_get)

        tipo = request.form["tipo"]
        elemento = request.form["elemento"]
        descripcion = request.form["descripcion"]
        tecnico = request.form["tecnico"]
        fecha = request.form["fecha"]
        estado = request.form["estado"]
        costo = request.form["costo"] or 0

        # Procesar foto opcional
        foto_url = None
        foto = request.files.get("foto")
        if foto and foto.filename:
            import os
            from PIL import Image
            from datetime import datetime as dt
            UPLOAD_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), "static", "uploads")
            os.makedirs(UPLOAD_DIR, exist_ok=True)
            nombre_foto = f"mant_{habitacion_id}_{dt.now().strftime('%Y%m%d_%H%M%S')}.webp"
            ruta_foto = os.path.join(UPLOAD_DIR, nombre_foto)
            try:
                img = Image.open(foto)
                if img.mode in ("RGBA", "P"):
                    img = img.convert("RGB")
                img.thumbnail((1080, 1080), Image.Resampling.LANCZOS)
                img.save(ruta_foto, "WEBP", quality=80)
                foto_url = f"/static/uploads/{nombre_foto}"
            except Exception:
                pass  # Guardar orden sin foto si falla

        deteccion_id = request.form.get("deteccion_id") or None
        if deteccion_id:
            try:
                deteccion_id = int(deteccion_id)
            except ValueError:
                deteccion_id = None

        conn = conectar()
        cursor = conn.cursor()

        cursor.execute("""
            INSERT INTO mantenimiento
            (habitacion_id, tipo, elemento, descripcion, tecnico, fecha, estado, costo, foto_url, deteccion_id)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
        """, (habitacion_id, tipo, elemento, descripcion, tecnico, fecha, estado, costo, foto_url, deteccion_id))

        # Sincronización automática de estado de la habitación
        nuevo_estado_hab = "En mantenimiento"
        if estado == "Completado":
            # Verificar si quedan otros mantenimientos pendientes para esta habitación
            cursor.execute("""
                SELECT COUNT(*) FROM mantenimiento 
                WHERE habitacion_id = %s AND estado != 'Completado'
            """, (habitacion_id,))
            pendientes_count = cursor.fetchone()[0]
            
            if pendientes_count == 0:
                nuevo_estado_hab = "Disponible"

        cursor.execute("""
            UPDATE habitaciones 
            SET estado = %s, fecha_ultimo_mantenimiento = CURRENT_DATE
            WHERE id = %s
        """, (nuevo_estado_hab, habitacion_id))

        conn.commit()
        cursor.close()
        conn.close()
        
        from extensions import socketio
        socketio.emit('alerta_global', {
            'tipo': 'mantenimiento',
            'habitacion_id': habitacion_id,
            'mensaje': f'Nueva orden de mantenimiento ({elemento})'
        })

        return redirect("/mantenimientos")

    deteccion_id = request.args.get("deteccion_id", "")
    return render_template("mantenimiento_form.html", habitacion_id=habitacion_id, deteccion_id=deteccion_id)


# ===============================
# Editar mantenimiento
# ===============================
@mantenimiento_bp.route("/mantenimiento/editar/<int:id>", methods=["GET", "POST"])
def editar(id):
    if "user" not in session:
        return redirect("/login")

    conn = conectar()
    cursor = conn.cursor()

    if request.method == "POST":
        tipo = request.form.get("tipo", "Correctivo")
        elemento = request.form.get("elemento", "Inspección detectada")
        descripcion = request.form.get("descripcion", "")
        tecnico = request.form.get("tecnico", "")
        fecha = request.form.get("fecha", "")
        estado = request.form.get("estado", "Pendiente")
        costo = request.form.get("costo", 0)
        prioridad = request.form.get("prioridad", "Media")

        cursor.execute("SELECT habitacion_id FROM mantenimiento WHERE id=%s", (id,))
        hab_row = cursor.fetchone()
        
        if hab_row:
            habitacion_id = hab_row[0]

            cursor.execute("""
                UPDATE mantenimiento
                SET tipo=%s, elemento=%s, descripcion=%s, tecnico=%s, fecha=%s, estado=%s, costo=%s, prioridad=%s
                WHERE id=%s
            """, (tipo, elemento, descripcion, tecnico, fecha, estado, costo, prioridad, id))

            # Sincronización automática de estado de la habitación
            nuevo_estado_hab = "En mantenimiento"
            if estado == "Completado":
                # Verificar si quedan otros mantenimientos pendientes para esta habitación
                cursor.execute("""
                    SELECT COUNT(*) FROM mantenimiento 
                    WHERE habitacion_id = %s AND estado != 'Completado'
                """, (habitacion_id,))
                pendientes_count = cursor.fetchone()[0]
                
                if pendientes_count == 0:
                    nuevo_estado_hab = "Disponible"

            cursor.execute("""
                UPDATE habitaciones 
                SET estado = %s, fecha_ultimo_mantenimiento = CURRENT_DATE
                WHERE id = %s
            """, (nuevo_estado_hab, habitacion_id))

            conn.commit()
            
        cursor.close()
        conn.close()
        
        from extensions import socketio
        socketio.emit('alerta_global', {
            'tipo': 'mantenimiento',
            'habitacion_id': habitacion_id,
            'mensaje': f'Orden de mantenimiento actualizada ({estado})'
        })
        
        # Redirigir dependiendo del rol (si es empleado mandarlo a su cola)
        if session.get("rol") == "empleado":
            return redirect("/mis_tareas")
        return redirect("/mantenimientos")

    cursor.execute("SELECT id, tipo, elemento, descripcion, tecnico, fecha, estado, costo, prioridad FROM mantenimiento WHERE id=%s", (id,))
    mantenimiento = cursor.fetchone()
    cursor.close()
    conn.close()

    if not mantenimiento:
        return redirect("/mantenimientos")

    return render_template("mantenimiento_editar.html", mantenimiento=mantenimiento)


# ===============================
# Cierre rápido de mantenimiento
# ===============================
@mantenimiento_bp.route("/mantenimiento/cerrar/<int:id>", methods=["POST"])
def cerrar(id):
    if "user" not in session:
        return redirect("/login")

    conn = conectar()
    cursor = conn.cursor()
    cursor.execute("SELECT habitacion_id FROM mantenimiento WHERE id=%s", (id,))
    row = cursor.fetchone()
    if row:
        habitacion_id = row[0]
        cursor.execute("UPDATE mantenimiento SET estado='Completado' WHERE id=%s", (id,))
        cursor.execute(
            "SELECT COUNT(*) FROM mantenimiento WHERE habitacion_id=%s AND estado != 'Completado'",
            (habitacion_id,)
        )
        if cursor.fetchone()[0] == 0:
            cursor.execute("""
                UPDATE habitaciones SET estado='Disponible', fecha_ultimo_mantenimiento=CURRENT_DATE
                WHERE id=%s
            """, (habitacion_id,))
        conn.commit()
    cursor.close()
    conn.close()
    return redirect(request.referrer or "/mantenimientos")


# ===============================
# Exportar lista de mantenimientos
# ===============================
@mantenimiento_bp.route("/mantenimientos/exportar")
def exportar():
    if "user" not in session:
        return redirect("/login")

    conn = conectar()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT h.numero, m.id, m.fecha, m.tipo, m.elemento,
               m.descripcion, m.tecnico, m.estado, m.costo
        FROM mantenimiento m
        JOIN habitaciones h ON m.habitacion_id = h.id
        ORDER BY m.fecha DESC
    """)
    rows = cursor.fetchall()
    cursor.close()
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Mantenimientos"

    encabezados = ["Habitación", "ID", "Fecha", "Tipo", "Elemento", "Descripción", "Técnico", "Estado", "Costo ($)"]
    header_fill = PatternFill("solid", fgColor="1E3A5F")
    header_font = Font(color="FFFFFF", bold=True)

    for col, titulo in enumerate(encabezados, start=1):
        cell = ws.cell(row=1, column=col, value=titulo)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for row in rows:
        ws.append([
            str(row[0]),
            int(row[1]),
            str(row[2]) if row[2] else "",
            row[3],
            row[4],
            row[5] or "",
            row[6] or "",
            row[7] or "",
            float(row[8]) if row[8] else 0.0,
        ])

    anchos = [14, 8, 14, 14, 25, 35, 20, 14, 12]
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = ancho

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nombre = f"mantenimientos_{date.today()}.xlsx"
    return send_file(buffer, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True, download_name=nombre)
